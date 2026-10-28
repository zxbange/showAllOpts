"""Read the officially disclosed market average guarantee ratio from CSData reports.

The list title and workbook header define the observation date. Publication dates,
file-directory dates, collateral totals and financing balances are not substitutes.
"""
from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from io import BytesIO, StringIO
import math
import re
import time
from urllib.parse import unquote, urljoin, urlsplit
import warnings
import zipfile

import httpx
import openpyxl
import xlrd


ORIGIN = 'https://www.csdata.cn'
LIST_PATHS = (
    '/cmsmc/cgfb/rzrqsj/rzrqscmrsjtj/',
    '/cmsmc/cgfb/rzrqsj/rzrqsclssj/rzrqscmrsjtj/',
)
INDEX_URL = ORIGIN + LIST_PATHS[0] + 'index.html'
HISTORY_URL = ORIGIN + LIST_PATHS[1] + 'index.html'
MAX_PAGES = 240
MAX_LIST_BYTES = 2 * 1024 * 1024
MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_UNPACKED_BYTES = 16 * 1024 * 1024
METRIC = '平均维持担保比例'
TITLE = '融资融券市场每日数据统计'
LIST_TITLES = (TITLE, '融资融券市场历史数据')


class MarginSourceError(ValueError):
    """The upstream document cannot safely be interpreted as an observation."""

    def __init__(self, message, *, cause_type=None):
        super().__init__(message)
        self.cause_type = cause_type


# Only these application-authored messages may reach persistent state or logs.
# Exception strings from HTTP clients, workbook parsers and callers are not safe.
_SAFE_REASONS = {
    '中证链接不在允许的 HTTPS 来源范围内。': 'unsafe_url',
    '中证链接路径无效。': 'unsafe_url',
    '中证链接不是预期的日报文件或目录。': 'unsafe_url',
    '中证服务器重定向缺少目标地址。': 'redirect_error',
    '中证服务器重定向次数过多。': 'redirect_error',
    '中证响应超过单次读取上限。': 'response_too_large',
    '中证响应读取超过单次时间上限。': 'response_timeout',
    '读取中证公开数据失败，请稍后重试。': 'network_error',
    '中证目录编码异常。': 'catalog_encoding',
    '中证日报标题缺少有效交易日期。': 'report_title_date',
    '中证目录分页指向了其他栏目。': 'catalog_pagination',
    '中证目录未找到日报，可能是网页结构变化；未将结果当作无新增。': 'catalog_empty',
    '中证目录分页循环或超过读取上限。': 'catalog_pagination',
    '请求日期范围内存在暂不支持的 PDF 日报，未将结果当作完整目录。': 'unsupported_pdf',
    '中证报表解压后过大。': 'workbook_too_large',
    '中证日报的行列范围异常。': 'workbook_dimensions',
    '中证返回的文件不是有效 XLS/XLSX 报表。': 'invalid_workbook',
    '中证报表解析失败。': 'workbook_parse',
    '平均维持担保比例不是有效数字。': 'invalid_ratio',
    '平均维持担保比例缺少明确的百分比单位，拒绝猜测倍率。': 'missing_percent_unit',
    '平均维持担保比例为空或不是有效数字。': 'invalid_ratio',
    '平均维持担保比例超出合理数据范围。': 'invalid_ratio',
    '平均维持担保比例缺失或同一行出现多个值。': 'ambiguous_ratio',
    '中证报表统计日期与日报标题不一致或缺失。': 'report_date_mismatch',
    '中证报表未提供唯一的平均维持担保比例。': 'ambiguous_ratio',
    '中证日报日期超出请求范围。': 'report_date_range',
    '中证日报日期与待采日期不一致。': 'report_date_mismatch',
    '中证日报比例不是有效数字或超出合理范围。': 'invalid_ratio',
}


def safe_error_details(exc):
    """Return bounded diagnostics without exposing arbitrary exception text."""
    name = type(exc).__name__
    error_type = name if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,79}', name) else 'Exception'
    details = dict(error_type=error_type, code='unexpected_error', reason='发生未预期错误，请稍后重试。')
    if type(exc) is MarginSourceError and len(exc.args) == 1 and isinstance(exc.args[0], str):
        reason = exc.args[0]
        if reason in _SAFE_REASONS:
            details.update(code=_SAFE_REASONS[reason], reason=reason)
        else:
            match = re.fullmatch(r'中证服务器返回 HTTP ([1-5][0-9]{2})。', reason)
            if match:
                details.update(code='http_status', reason=reason, http_status=int(match[1]))
    if (type(exc) is MarginSourceError and isinstance(exc.cause_type, str)
            and re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,79}', exc.cause_type)):
        details['cause_type'] = exc.cause_type
    return details


def safe_report_url(url):
    """Omit untrusted URLs from diagnostics, including credentials and queries."""
    if not isinstance(url, str) or len(url) > 2048 or any(ord(c) < 32 or ord(c) == 127 for c in url):
        return None
    try:
        return _safe_url(url, 'report')
    except (MarginSourceError, ValueError):
        return None


@dataclass(frozen=True)
class MarginReport:
    trade_date: date
    url: str


def _safe_url(url, kind):
    parts = urlsplit(url)
    if (parts.scheme != 'https' or parts.netloc not in ('www.csdata.cn', 'csdata.cn')
            or parts.query or parts.fragment):
        raise MarginSourceError('中证链接不在允许的 HTTPS 来源范围内。')
    path = unquote(parts.path)
    if '\\' in path or '..' in path.split('/'):
        raise MarginSourceError('中证链接路径无效。')
    if kind in ('report', 'catalog_report'):
        # Legacy PDFs are catalog metadata only, never accepted for downloading.
        extensions = ('.xlsx', '.xls', '.pdf') if kind == 'catalog_report' else ('.xlsx', '.xls')
        valid = path.startswith('/cmsmc/articleFileDir/') and path.lower().endswith(extensions)
    else:
        valid = any(path.startswith(prefix) and re.fullmatch(r'(?:index|[a-zA-Z0-9]+-\d+)\.html', path[len(prefix):])
                    for prefix in LIST_PATHS)
    if not valid:
        raise MarginSourceError('中证链接不是预期的日报文件或目录。')
    return url


def _download(client, url, kind):
    limit = MAX_FILE_BYTES if kind == 'report' else MAX_LIST_BYTES
    started = time.monotonic()
    try:
        for _ in range(4):
            _safe_url(url, kind)
            with client.stream('GET', url, follow_redirects=False, timeout=httpx.Timeout(25., connect=10.)) as response:
                if response.status_code in (301, 302, 303, 307, 308):
                    location = response.headers.get('location')
                    if not location:
                        raise MarginSourceError('中证服务器重定向缺少目标地址。')
                    url = urljoin(url, location)
                    continue
                response.raise_for_status()
                length = response.headers.get('content-length', '')
                if length.isdigit() and int(length) > limit:
                    raise MarginSourceError('中证响应超过单次读取上限。')
                chunks, size = [], 0
                for chunk in response.iter_bytes():
                    if time.monotonic() - started > 60:
                        raise MarginSourceError('中证响应读取超过单次时间上限。')
                    size += len(chunk)
                    if size > limit:
                        raise MarginSourceError('中证响应超过单次读取上限。')
                    chunks.append(chunk)
                return b''.join(chunks)
        raise MarginSourceError('中证服务器重定向次数过多。')
    except httpx.HTTPStatusError as exc:
        raise MarginSourceError(f'中证服务器返回 HTTP {exc.response.status_code}。',
                                cause_type=type(exc).__name__) from None
    except httpx.HTTPError as exc:
        raise MarginSourceError('读取中证公开数据失败，请稍后重试。',
                                cause_type=type(exc).__name__) from None


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    match = re.search(r'(?<!\d)(20\d{2})(?:年|[-/])?(\d{1,2})(?:月|[-/])?(\d{1,2})(?:日)?(?!\d)', value.strip())
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            return None
    return None


class _Links(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links, self.current = [], None

    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            self.current = (dict(attrs), [])

    def handle_data(self, data):
        if self.current is not None:
            self.current[1].append(data)

    def handle_endtag(self, tag):
        if tag == 'a' and self.current is not None:
            attrs, pieces = self.current
            self.links.append((attrs, ''.join(pieces).strip()))
            self.current = None


def parse_listing(content, url):
    """Only follow the site's own next-page link; never guess page identifiers."""
    parser = _Links()
    try:
        parser.feed(content.decode('utf-8-sig'))
    except UnicodeError:
        raise MarginSourceError('中证目录编码异常。') from None
    reports, next_url = [], None
    for attrs, text in parser.links:
        title = attrs.get('title') or text
        href = attrs.get('href')
        if (any(name in title for name in LIST_TITLES) and href
                and urlsplit(href).path.lower().endswith(('.xls', '.xlsx', '.pdf'))):
            trade_date = _parse_date(title)
            if trade_date is None:
                raise MarginSourceError('中证日报标题缺少有效交易日期。')
            reports.append(MarginReport(trade_date, _safe_url(urljoin(url, href), 'catalog_report')))
        if re.sub(r'\s+', '', text) in ('下页', '下一页'):
            target = href
            if not target:
                match = re.search(r"queryArticleByCondition\(\s*this\s*,\s*['\"]([^'\"]+)['\"]\s*\)", attrs.get('onclick', ''))
                target = match.group(1) if match else None
            if target:
                next_url = _safe_url(urljoin(url, target), 'list')
                if urlsplit(next_url).path.rsplit('/', 1)[0] != urlsplit(url).path.rsplit('/', 1)[0]:
                    raise MarginSourceError('中证目录分页指向了其他栏目。')
    if not reports:
        raise MarginSourceError('中证目录未找到日报，可能是网页结构变化；未将结果当作无新增。')
    return reports, next_url


def list_reports(client, start: date, end: date, check_cancel=lambda: None) -> list[MarginReport]:
    if start > end:
        raise ValueError('开始日期不能晚于结束日期。')
    results, visited = {}, set()
    for index in (INDEX_URL, HISTORY_URL):
        url = index
        while url:
            check_cancel()
            if url in visited or len(visited) >= MAX_PAGES:
                raise MarginSourceError('中证目录分页循环或超过读取上限。')
            visited.add(url)
            reports, next_url = parse_listing(_download(client, url, 'list'), url)
            for item in reports:
                if start <= item.trade_date <= end:
                    if urlsplit(item.url).path.lower().endswith('.pdf'):
                        raise MarginSourceError('请求日期范围内存在暂不支持的 PDF 日报，未将结果当作完整目录。')
                    # Current-column records take precedence over archive duplicates.
                    results.setdefault(item.trade_date, item)
            # Publication/maintenance order is not observation-date order. Even
            # an entirely old page can precede in-range reports on a later page.
            # Follow the complete bounded catalog rather than silently omit days.
            url = next_url
    return sorted(results.values(), key=lambda item: item.trade_date)


def _workbook_rows(content):
    """Yield nonempty cells with their number format; cap decompression and dimensions."""
    try:
        if content.startswith(b'PK\x03\x04'):
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 1000 or sum(item.file_size for item in members) > MAX_UNPACKED_BYTES:
                    raise MarginSourceError('中证报表解压后过大。')
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore', message='Workbook contains no default style.*', category=UserWarning)
                book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    if (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 30:
                        raise MarginSourceError('中证日报的行列范围异常。')
                    for row in sheet.iter_rows(max_row=sheet.max_row or 500, max_col=sheet.max_column or 30):
                        yield [(cell.value, cell.number_format) for cell in row if cell.value not in (None, '')]
            finally:
                book.close()
        elif content.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
            book = xlrd.open_workbook(file_contents=content, formatting_info=True, logfile=StringIO())
            try:
                for sheet in book.sheets():
                    if sheet.nrows > 500 or sheet.ncols > 30:
                        raise MarginSourceError('中证日报的行列范围异常。')
                    for i in range(sheet.nrows):
                        row = []
                        for cell in sheet.row(i):
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, book.datemode)
                            if value not in (None, ''):
                                fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                                row.append((value, fmt))
                        yield row
            finally:
                book.release_resources()
        else:
            raise MarginSourceError('中证返回的文件不是有效 XLS/XLSX 报表。')
    except MarginSourceError:
        raise
    except Exception as exc:
        raise MarginSourceError('中证报表解析失败。', cause_type=type(exc).__name__) from None


def _ratio(value, number_format, explicit_percent=False):
    if isinstance(value, bool):
        raise MarginSourceError('平均维持担保比例不是有效数字。')
    try:
        if isinstance(value, str) and value.strip().endswith(('%', '％')):
            ratio = float(value.strip()[:-1].replace(',', ''))
        else:
            numeric = float(value)
            # Literal/escaped percent signs do not scale Excel numbers.
            actual_format = re.sub(r'"[^"]*"|\\.', '', number_format)
            if '%' in actual_format:
                ratio = numeric * 100
            elif explicit_percent:
                ratio = numeric
            else:
                raise MarginSourceError('平均维持担保比例缺少明确的百分比单位，拒绝猜测倍率。')
    except (TypeError, ValueError) as exc:
        if isinstance(exc, MarginSourceError):
            raise
        raise MarginSourceError('平均维持担保比例为空或不是有效数字。') from None
    if not math.isfinite(ratio) or ratio <= 0 or ratio > 10000:
        raise MarginSourceError('平均维持担保比例超出合理数据范围。')
    return round(ratio, 6)


def parse_report(content, expected_date: date):
    dates, values = set(), []
    for row in _workbook_rows(content):
        labels = [re.sub(r'\s+', '', str(value)) for value, _ in row]
        if '统计指标' in labels:
            dates.update(day for value, _ in row if (day := _parse_date(value)) is not None)
        for i, label in enumerate(labels):
            if label not in (METRIC, METRIC + '(%)', METRIC + '（%）', METRIC + '(％)', METRIC + '（％）'):
                continue
            if len(row[i+1:]) != 1:
                raise MarginSourceError('平均维持担保比例缺失或同一行出现多个值。')
            values.append(_ratio(*row[i+1], explicit_percent=label != METRIC))
    if dates != {expected_date}:
        raise MarginSourceError('中证报表统计日期与日报标题不一致或缺失。')
    if len(values) != 1:
        raise MarginSourceError('中证报表未提供唯一的平均维持担保比例。')
    return values[0]


def fetch_point(client, report: MarginReport):
    content = _download(client, report.url, 'report')
    ratio = parse_report(content, report.trade_date)
    return {'date': report.trade_date.isoformat(), 'ratio': ratio, 'source_url': report.url}
