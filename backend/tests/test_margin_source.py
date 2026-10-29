"""Official XLS/XLSX layouts and bounded HTTP behavior; never fetch live data."""
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import httpx
import openpyxl
import pytest

from app.services import margin_source as source


DAY = date(2026, 10, 26)
REPORT_URL = source.ORIGIN + '/cmsmc/articleFileDir/2026-10/26/sample.xlsx'


def workbook(value=2.556273, fmt='#,##0.##%', day=DAY, label=source.METRIC):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A1'], sheet['A2'] = source.TITLE, '统计指标'
    sheet['C2'] = datetime.combine(day, datetime.min.time()) if day else '日期缺失'
    sheet['B19'], sheet['C19'] = '担保物总价值（亿元）', 72774.01291568
    sheet['B20'], sheet['C20'] = label, value
    sheet['C20'].number_format = fmt
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def listing(dates, next_path=None, extension='xlsx'):
    links = ['<li><span>2099-01-01</span><a title="{}{}" href="/cmsmc/articleFileDir/sample-{}.{}">{}</a></li>'.format(
        source.TITLE, day.strftime('%Y%m%d'), day.isoformat(), extension, source.TITLE) for day in dates]
    if next_path:
        links.append('<a onclick="queryArticleByCondition(this,\'{}\')">下页</a>'.format(next_path))
    return ('<html>' + ''.join(links) + '</html>').encode()


def test_current_xlsx_uses_percent_format_and_report_date():
    assert source.parse_report(workbook(), DAY) == pytest.approx(255.6273)


def test_old_report_text_percent_is_not_multiplied_twice():
    assert source.parse_report(workbook('272.9%', '0.0%', date(2020, 10, 27)), date(2020, 10, 27)) == 272.9


def test_numeric_percentage_with_explicit_label_unit():
    assert source.parse_report(workbook(272.9, 'General', label=source.METRIC+'（%）'), DAY) == 272.9


@pytest.mark.parametrize('value,fmt', [(2.55, 'General'), ('--', '0.0%'), (None, '0.0%'), (0, '0.0%'),
                                       (True, '0.0%'), ('NaN%', '0.0%'), (2.55, '0.0"%"')])
def test_unknown_units_and_missing_values_are_errors(value, fmt):
    with pytest.raises(source.MarginSourceError):
        source.parse_report(workbook(value, fmt), DAY)


@pytest.mark.parametrize('day', [date(2026, 10, 23), None])
def test_date_mismatch_or_missing_never_becomes_an_observation(day):
    with pytest.raises(source.MarginSourceError, match='日期'):
        source.parse_report(workbook(day=day), DAY)


def test_xls_reader_preserves_legacy_string_values_and_number_format(monkeypatch):
    # The official 2020-10-27 XLS has a string date in D2 and '272.9%' in D20.
    def cell(value, ctype=1):
        return SimpleNamespace(value=value, ctype=ctype, xf_index=0)
    rows = [[cell('统计指标'), cell(''), cell(''), cell('2020年10月27日')],
            [cell(''), cell(source.METRIC), cell(''), cell('272.9%')]]
    released = []
    sheet = SimpleNamespace(nrows=len(rows), ncols=4, row=lambda i: rows[i])
    book = SimpleNamespace(sheets=lambda: [sheet], xf_list=[SimpleNamespace(format_key=0)],
                           format_map={0: SimpleNamespace(format_str='0.0%')}, datemode=0,
                           release_resources=lambda: released.append(True))
    monkeypatch.setattr(source.xlrd, 'open_workbook', lambda **kwargs: book)
    assert source.parse_report(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1', date(2020, 10, 27)) == 272.9
    assert released == [True]


def test_missing_metric_and_invalid_file_are_not_zero():
    with pytest.raises(source.MarginSourceError, match='唯一'):
        source.parse_report(workbook(label='融资融券余额'), DAY)
    with pytest.raises(source.MarginSourceError, match='XLS'):
        source.parse_report(b'<html>request denied</html>', DAY)


def test_listing_uses_title_not_publication_date_and_real_next_link():
    next_path = source.LIST_PATHS[0]+'d2bb5c19-2.html'
    reports, next_url = source.parse_listing(listing([DAY], next_path), source.INDEX_URL)
    assert reports[0].trade_date == DAY
    assert next_url == source.ORIGIN + next_path


def test_list_reports_scans_past_old_pages_and_checks_cancellation():
    next_path = source.LIST_PATHS[0]+'d2bb5c19-2.html'
    third_path = source.LIST_PATHS[0]+'d2bb5c19-3.html'
    archive_path = source.LIST_PATHS[1]+'archive-2.html'
    visited, cancelled = [], []
    pages = {
        source.INDEX_URL: listing([date(2020, 1, 1), DAY], next_path),
        source.ORIGIN+next_path: listing([date(2019, 1, 1)], third_path),
        source.ORIGIN+third_path: listing([date(2026, 10, 23)]),
        source.HISTORY_URL: listing([date(2018, 1, 1)], archive_path),
        source.ORIGIN+archive_path: listing([date(2026, 10, 22), date(2026, 10, 27)]),
    }
    def response(request):
        visited.append(str(request.url))
        return httpx.Response(200, content=pages[str(request.url)])
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        reports = source.list_reports(client, date(2026, 10, 22), DAY, lambda: cancelled.append(True))
    assert [r.trade_date for r in reports] == [date(2026, 10, 22), date(2026, 10, 23), DAY]
    assert len(visited) == len(cancelled) == 5


def test_page_limit_rejects_partial_results_even_after_old_pages(monkeypatch):
    monkeypatch.setattr(source, 'MAX_PAGES', 2)
    next_path = source.LIST_PATHS[0]+'d2bb5c19-2.html'
    pages = {
        source.INDEX_URL: listing([DAY], next_path),
        source.ORIGIN+next_path: listing([date(2019, 1, 1)]),
    }
    visited = []
    def response(request):
        visited.append(str(request.url))
        return httpx.Response(200, content=pages[str(request.url)])
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        with pytest.raises(source.MarginSourceError, match='上限'):
            source.list_reports(client, DAY, DAY)
    assert len(visited) == 2


def test_cancellation_during_catalog_scan_stops_before_next_request():
    visited, checks = [], []
    def cancel():
        checks.append(True)
        if len(checks) == 2:
            raise RuntimeError('cancelled')
    def response(request):
        visited.append(str(request.url))
        return httpx.Response(200, content=listing([date(2019, 1, 1)],
                              source.LIST_PATHS[0]+'d2bb5c19-2.html'))
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        with pytest.raises(RuntimeError, match='cancelled'):
            source.list_reports(client, DAY, DAY, cancel)
    assert visited == [source.INDEX_URL]


def test_duplicate_archive_day_keeps_current_report():
    def response(request):
        suffix = 'xlsx' if str(request.url) == source.INDEX_URL else 'xls'
        return httpx.Response(200, content=listing([DAY], extension=suffix))
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        reports = source.list_reports(client, DAY, DAY)
    assert len(reports) == 1 and reports[0].url.endswith('.xlsx')


def test_cyclic_pagination_is_rejected():
    def response(request):
        return httpx.Response(200, content=listing([DAY], source.LIST_PATHS[0]+'index.html'))
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        with pytest.raises(source.MarginSourceError, match='循环'):
            source.list_reports(client, DAY, DAY)


def test_empty_changed_catalog_does_not_mean_no_new_data():
    with pytest.raises(source.MarginSourceError, match='目录未找到'):
        source.parse_listing(b'<html>unavailable</html>', source.INDEX_URL)


@pytest.mark.parametrize('dates', [
    [date(2026, 10, 23), DAY, date(2026, 10, 22)],
    # Observed current page 15 and archive page 3 on 2026-10-28.
    [date(2026, 3, 30), date(2026, 3, 26), date(2026, 3, 27)],
    [date(2024, 12, 3), date(2024, 11, 29), date(2024, 12, 2)],
])
def test_listing_accepts_non_monotonic_trade_dates(dates):
    reports, next_url = source.parse_listing(listing(dates), source.INDEX_URL)
    assert [r.trade_date for r in reports] == dates
    assert next_url is None


@pytest.mark.parametrize('url', [
    'http://www.csdata.cn/cmsmc/articleFileDir/sample.xls',
    'https://www.csdata.cn.evil.invalid/cmsmc/articleFileDir/sample.xls',
    'https://www.csdata.cn@evil.invalid/cmsmc/articleFileDir/sample.xls',
    'https://www.csdata.cn/cmsmc/articleFileDir/%2e%2e/secrets.xls',
    'https://www.csdata.cn/cmsmc/articleFileDir/a.xls?redirect=evil',
    'https://www.csdata.cn/other/a.xls',
])
def test_source_whitelist_rejects_unsafe_urls_before_any_request(url):
    with httpx.Client(transport=httpx.MockTransport(lambda req: pytest.fail('unsafe request'))) as client:
        with pytest.raises(source.MarginSourceError):
            source.fetch_point(client, source.MarginReport(DAY, url))


def test_cross_domain_redirect_is_not_followed():
    visited = []
    def response(request):
        visited.append(str(request.url))
        return httpx.Response(302, headers={'location': 'https://example.invalid/data.xlsx'})
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        with pytest.raises(source.MarginSourceError):
            source.fetch_point(client, source.MarginReport(DAY, REPORT_URL))
    assert visited == [REPORT_URL]


def test_fetch_point_returns_percentage_and_provenance():
    with httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(200, content=workbook()))) as client:
        point = source.fetch_point(client, source.MarginReport(DAY, REPORT_URL))
    assert point == dict(date='2026-10-26', ratio=255.6273, source_url=REPORT_URL)


def test_http_errors_and_oversized_responses_are_explicit():
    for response in [httpx.Response(503), httpx.Response(200, headers={'content-length': str(source.MAX_FILE_BYTES+1)})]:
        with httpx.Client(transport=httpx.MockTransport(lambda request: response)) as client:
            with pytest.raises(source.MarginSourceError):
                source.fetch_point(client, source.MarginReport(DAY, REPORT_URL))


def test_cancellation_happens_before_network_request():
    def cancel():
        raise RuntimeError('cancelled')
    with httpx.Client(transport=httpx.MockTransport(lambda request: pytest.fail('unexpected request'))) as client:
        with pytest.raises(RuntimeError, match='cancelled'):
            source.list_reports(client, DAY, DAY, cancel)


def test_old_pdf_catalog_page_is_traversed_without_downloading_pdf():
    next_path = source.LIST_PATHS[1]+'archive-2.html'
    visited = []
    pages = {
        source.INDEX_URL: listing([DAY]),
        source.HISTORY_URL: listing([date(2019, 7, 8)], next_path, extension='pdf'),
        source.ORIGIN+next_path: listing([date(2026, 10, 23)]),
    }
    def response(request):
        visited.append(str(request.url))
        return httpx.Response(200, content=pages[str(request.url)])
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        reports = source.list_reports(client, date(2026, 10, 23), DAY)
    assert [r.trade_date for r in reports] == [date(2026, 10, 23), DAY]
    assert visited == list(pages)


def test_in_range_pdf_is_explicit_error_not_silently_missing():
    with httpx.Client(transport=httpx.MockTransport(
            lambda request: httpx.Response(200, content=listing([DAY], extension='pdf')))) as client:
        with pytest.raises(source.MarginSourceError, match='PDF'):
            source.list_reports(client, DAY, DAY)


def test_pdf_metadata_does_not_expand_download_allowlist():
    url = source.ORIGIN+'/cmsmc/articleFileDir/old.pdf'
    with httpx.Client(transport=httpx.MockTransport(lambda request: pytest.fail('unexpected PDF download'))) as client:
        with pytest.raises(source.MarginSourceError):
            source.fetch_point(client, source.MarginReport(DAY, url))


def test_legacy_pdf_metadata_still_requires_safe_url_and_valid_date():
    for content in [
        listing([DAY], extension='pdf').replace(b'/cmsmc/articleFileDir/', b'https://evil.invalid/'),
        listing([DAY], extension='pdf').replace(b'20261026', b'20261399'),
    ]:
        with pytest.raises(source.MarginSourceError):
            source.parse_listing(content, source.HISTORY_URL)


def test_legacy_history_title_is_recognized_without_changing_date_semantics():
    # Archive page 127 uses this title for 2015 XLS disclosures.
    day = date(2015, 4, 1)
    content = listing([day], extension='xls').replace(
        source.TITLE.encode(), '融资融券市场历史数据'.encode())
    reports, next_url = source.parse_listing(content, source.HISTORY_URL)
    assert [r.trade_date for r in reports] == [day]
    assert next_url is None


def test_trading_day_gaps_do_not_create_calendar_day_observations():
    days = [date(2026, 10, 26), date(2026, 10, 23), date(2026, 10, 9)]
    with httpx.Client(transport=httpx.MockTransport(
            lambda request: httpx.Response(200, content=listing(days)))) as client:
        reports = source.list_reports(client, date(2026, 10, 9), date(2026, 10, 26))
    assert [r.trade_date for r in reports] == sorted(days)


def test_unsupported_catalog_extension_is_not_a_valid_empty_page():
    with pytest.raises(source.MarginSourceError, match='目录未找到'):
        source.parse_listing(listing([DAY], extension='zip'), source.HISTORY_URL)


def numeric_header_workbook(serial, epoch, *, extra_date=None, header='统计指标'):
    book = openpyxl.load_workbook(BytesIO(workbook()))
    book.epoch = epoch
    sheet = book.active
    sheet['A2'], sheet['C2'] = header, serial
    sheet['C2'].number_format = 'General'
    if extra_date is not None:
        sheet['D2'] = extra_date
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


@pytest.mark.parametrize('epoch', [openpyxl.utils.datetime.CALENDAR_WINDOWS_1900,
                                   openpyxl.utils.datetime.CALENDAR_MAC_1904])
def test_general_numeric_header_uses_xlsx_epoch(epoch):
    serial = openpyxl.utils.datetime.to_excel(datetime.combine(DAY, datetime.min.time()), epoch)
    assert source.parse_report(numeric_header_workbook(serial, epoch), DAY) == 255.6273
    with pytest.raises(source.MarginSourceError, match='日期'):
        source.parse_report(numeric_header_workbook(serial + 1, epoch), DAY)
    with pytest.raises(source.MarginSourceError, match='日期'):
        source.parse_report(numeric_header_workbook(serial, epoch, extra_date='2026-10-23'), DAY)
    with pytest.raises(source.MarginSourceError, match='日期'):
        source.parse_report(numeric_header_workbook(serial, epoch, header='非统计表头'), DAY)


@pytest.mark.parametrize('serial', [True, False, float('nan'), float('inf'), -float('inf'),
                                   -1, 0, 0.5, 45804.5, 2958466, 10**100])
def test_invalid_numeric_header_dates_are_not_observations(serial):
    epoch = openpyxl.utils.datetime.CALENDAR_WINDOWS_1900
    assert source._report_header_date(serial, epoch) is None


@pytest.mark.parametrize('datemode', [0, 1])
@pytest.mark.parametrize('date_ctype', [source.xlrd.XL_CELL_NUMBER, source.xlrd.XL_CELL_DATE])
def test_xls_numeric_header_uses_workbook_datemode(monkeypatch, datemode, date_ctype):
    epoch = (openpyxl.utils.datetime.CALENDAR_MAC_1904 if datemode
             else openpyxl.utils.datetime.CALENDAR_WINDOWS_1900)
    serial = openpyxl.utils.datetime.to_excel(datetime.combine(DAY, datetime.min.time()), epoch)
    def cell(value, ctype=source.xlrd.XL_CELL_TEXT):
        return SimpleNamespace(value=value, ctype=ctype, xf_index=0)
    rows = [[cell('统计指标'), cell(serial, date_ctype)],
            [cell(source.METRIC + '（%）'), cell(264.1024, source.xlrd.XL_CELL_NUMBER)]]
    released = []
    sheet = SimpleNamespace(nrows=len(rows), ncols=2, row=lambda i: rows[i])
    book = SimpleNamespace(sheets=lambda: [sheet], xf_list=[SimpleNamespace(format_key=0)],
                           format_map={0: SimpleNamespace(format_str='General')}, datemode=datemode,
                           release_resources=lambda: released.append(True))
    monkeypatch.setattr(source.xlrd, 'open_workbook', lambda **kwargs: book)
    assert source.parse_report(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1', DAY) == 264.1024
    assert released == [True]


@pytest.mark.parametrize('serial,day,ratio', [(45800, date(2025, 5, 23), 263.6191),
                                           (45803, date(2025, 5, 26), 264.7351),
                                           (45804, date(2025, 5, 27), 264.1024)])
def test_official_may_2025_layout_generated_regression(serial, day, ratio):
    # Reproduce the public reports' General-formatted C2; no downloaded fixtures.
    book = openpyxl.load_workbook(BytesIO(numeric_header_workbook(
        serial, openpyxl.utils.datetime.CALENDAR_WINDOWS_1900)))
    book.active['C20'] = ratio / 100
    out = BytesIO()
    book.save(out)
    book.close()
    assert source.parse_report(out.getvalue(), day) == ratio


def test_pinned_title_correction_keeps_strict_workbook_validation(monkeypatch):
    correction = source.TITLE_DATE_CORRECTION
    actual = date.fromisoformat(correction['observation_date'])
    report = source.MarginReport(date.fromisoformat(correction['report_date']), correction['source_url'])
    # Generated layout exercises the gate independently of the downloaded file.
    content = workbook('261.9%', 'General', actual)
    real_sha256 = source.hashlib.sha256
    monkeypatch.setattr(source.hashlib, 'sha256', lambda data: SimpleNamespace(
        hexdigest=lambda: correction['content_sha256']))
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(200, content=content))) as client:
        point = source.fetch_point(client, report)
        assert point == dict(date='2023-09-28', ratio=261.9, source_url=report.url,
                             source_correction=correction)
        for other in [source.MarginReport(report.trade_date, REPORT_URL),
                      source.MarginReport(date(2023, 10, 5), report.url)]:
            with pytest.raises(source.MarginSourceError, match='日期'):
                source.fetch_point(client, other)
    for bad in [workbook(day=report.trade_date), workbook(day=None),
                workbook(day=actual, label='融资余额'), workbook(261.9, 'General', actual)]:
        with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(200, content=bad))) as client:
            with pytest.raises(source.MarginSourceError):
                source.fetch_point(client, report)
    monkeypatch.setattr(source.hashlib, 'sha256', real_sha256)
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(200, content=content))) as client:
        with pytest.raises(source.MarginSourceError, match='日期'):
            source.fetch_point(client, report)
